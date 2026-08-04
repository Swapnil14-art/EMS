"""
Auto-generates a structured .docx RnD post-event report.
Uses report_template.docx as a template to inherit header (dual logos) and footer (Page X of Y).
Called after RnD report submission — saves file to /events/{event_id}/rnd_report/generated_rnd_report.docx
"""
import os
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from datetime import datetime
from PIL import Image as PILImage

from app.models.event import Event
from app.models.event_rnd_report import EventRndReport
from app.config import settings

# Path to the template docx that has the correct header/footer
TEMPLATE_DOCX = os.path.join(os.path.dirname(__file__), "..", "templates", "report_template.docx")


def _add_bold_label_value(doc: Document, label: str, value: str):
    """Add a paragraph with bold label followed by normal value."""
    p = doc.add_paragraph()
    run_label = p.add_run(label)
    run_label.bold = True
    p.add_run(value or "N/A")
    return p


def _add_section_heading(doc: Document, text: str):
    """Add a styled section heading."""
    p = doc.add_paragraph()
    p.style = doc.styles["Heading 2"]
    run = p.add_run(text)
    return p


def _add_subsection_heading(doc: Document, text: str):
    p = doc.add_paragraph()
    p.style = doc.styles["Heading 3"]
    p.add_run(text)
    return p


def _table_row(table, label: str, value: str):
    row = table.add_row()
    row.cells[0].text = label
    row.cells[1].text = str(value) if value is not None else "N/A"
    row.cells[0].paragraphs[0].runs[0].bold = True


def generate_rnd_report(event: Event, report: EventRndReport) -> str:
    """
    Generate a .docx RnD report for the event using the institutional template
    (which carries the dual-logo header and page-number footer).
    Returns the path where the file was saved.
    """
    # Determine template path — fall back gracefully if not deployed yet
    template_path = TEMPLATE_DOCX
    if not os.path.exists(template_path):
        base = os.path.dirname(os.path.abspath(__file__))
        candidates = [
            os.path.join(base, "..", "templates", "report_template.docx"),
            os.path.join(base, "report_template.docx"),
            os.path.join(settings.STORAGE_ROOT, "report_template.docx"),
        ]
        for c in candidates:
            if os.path.exists(c):
                template_path = c
                break
        else:
            template_path = None

    # Open template (inherits header/footer) or blank document
    if template_path:
        doc = Document(template_path)
        # Clear all body content from template, keep header/footer
        body = doc.element.body
        for child in list(body):
            tag = child.tag.split("}")[-1] if "}" in child.tag else child.tag
            if tag not in ("sectPr",):
                body.remove(child)
    else:
        doc = Document()

    # ── PAGE TITLE ──────────────────────────────────────────────────────────
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("RnD Report on")
    run.bold = True
    run.font.size = Pt(14)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    event_title_text = f'"{event.title}"'
    run2 = subtitle.add_run(event_title_text)
    run2.bold = True
    run2.font.size = Pt(13)

    doc.add_paragraph()  # spacer

    preserve_ratio = False
    if report.social_pamphlet:
        if isinstance(report.social_pamphlet, dict):
            preserve_ratio = report.social_pamphlet.get("preserve_aspect_ratio", False)
        else:
            preserve_ratio = getattr(report.social_pamphlet, "preserve_aspect_ratio", False)

    # ── GUEST SPEAKER(S) ──────────────────────────────────────────────────────
    speakers = report.guest_speakers or []
    if speakers:
        p_hdr = doc.add_paragraph()
        p_hdr.add_run("Name and designation of the Guest Speakers/ Judges/ Mentors etc.:").bold = True
        for i, spk in enumerate(speakers, 1):
            p = doc.add_paragraph()
            spk_type = spk.get("speaker_type") or "Guest Speaker"
            if spk_type == "Others":
                spk_type = spk.get("custom_speaker_type") or "Others"
            p.add_run(f"{spk_type}:\n").bold = True
            p.add_run("Name: ").bold = True
            p.add_run(f"{spk.get('name', 'N/A')}\n")
            p.add_run("Designation: ").bold = True
            p.add_run(f"{spk.get('designation', 'N/A')}\n")
            p.add_run("Organization: ").bold = True
            p.add_run(f"{spk.get('organization', 'N/A')}\n")
            p.add_run("Area of Expertise: ").bold = True
            p.add_run(f"{spk.get('expertise', 'N/A')}")
    else:
        _add_bold_label_value(doc, "Name and designation of the Guest Speakers/ Judges/ Mentors etc.: ", "N/A")
    doc.add_paragraph()

    # ── VENUE & DATES ────────────────────────────────────────────────────────
    _add_bold_label_value(doc, "Venue: ", event.venue_custom or (event.venue.name if event.venue else "N/A"))
    _add_bold_label_value(doc, "Start Date: ", event.start_datetime.strftime("%d %B %Y"))
    _add_bold_label_value(doc, "End Date: ", event.end_datetime.strftime("%d %B %Y"))
    _add_bold_label_value(doc, "Time: ", f"{event.start_datetime.strftime('%I:%M %p')} – {event.end_datetime.strftime('%I:%M %p')}")

    # Duration in hours
    duration_hrs = round(
        (event.end_datetime - event.start_datetime).total_seconds() / 3600, 2
    )
    _add_bold_label_value(doc, "Duration (hrs): ", str(duration_hrs))
    _add_bold_label_value(doc, "Department: ", event.school_department or "N/A")
    _add_bold_label_value(doc, "Organizing Club: ", event.club.name if event.club else "Non-Club Event")
    doc.add_paragraph()

    # ── SOCIAL MEDIA LINKS ────────────────────────────────────────────────────
    def add_social_links(label, links_dict):
        p = doc.add_paragraph()
        p.add_run(label).bold = True
        if links_dict:
            has_any = False
            for platform in ("linkedin", "facebook", "instagram", "x"):
                val = links_dict.get(platform, "")
                if val:
                    has_any = True
                    bp = doc.add_paragraph(style="List Bullet")
                    bp.add_run(f"{platform.title()}: ").bold = True
                    bp.add_run(val)
            if not has_any:
                doc.add_paragraph("  N/A")
        else:
            doc.add_paragraph("  N/A")

    add_social_links("Link of Social Media Post of E-Pamphlet:", report.social_pamphlet)
    add_social_links("Link of Social Media Post of Video:", report.social_video)
    doc.add_paragraph()

    # ── PROGRAM TYPE ──────────────────────────────────────────────────────────
    _add_bold_label_value(doc, "Program Type: ", report.program_type or "N/A")
    doc.add_paragraph()

    # ── OBJECTIVE & LEARNING ──────────────────────────────────────────────────
    _add_bold_label_value(doc, "Objective of the Activity (100 chars): ", report.objective or "N/A")
    _add_bold_label_value(doc, "Benefit in Terms of Learning / Skills / Knowledge (150 chars): ", report.learning_benefit or "N/A")
    doc.add_paragraph()

    # ── COORDINATORS ──────────────────────────────────────────────────────────
    faculty = report.faculty_coordinators or []
    students = report.student_coordinators or []
    _add_bold_label_value(doc, "Faculty Coordinators: ", ", ".join(faculty) if faculty else "N/A")
    _add_bold_label_value(doc, "Student Coordinators: ", ", ".join(students) if students else "N/A")
    doc.add_paragraph()

    # ── PARTICIPANTS & BUDGET ─────────────────────────────────────────────────
    _add_bold_label_value(doc, "Number of Student Participants: ", str(report.student_count or 0))
    _add_bold_label_value(doc, "Number of Faculty Participants: ", str(report.faculty_count or 0))
    _add_bold_label_value(doc, "Number of External Participants: ", str(report.external_count or 0))
    _add_bold_label_value(doc, "Total Participants: ", str(report.participant_count or 0))
    _add_bold_label_value(doc, "Estimated Budget: ", f"Rs. {event.budget:,.2f}" if event.budget else "N/A")
    _add_bold_label_value(doc, "Actual Expenditure: ", f"Rs. {report.actual_budget:,.2f}" if report.actual_budget else "N/A")
    doc.add_paragraph()

    # ── MODE OF DELIVERY ──────────────────────────────────────────────────────
    _add_bold_label_value(doc, "Mode of Delivery: ", (report.mode_of_delivery or "N/A").title())
    doc.add_paragraph()

    # ── SPEAKER BACKGROUND ────────────────────────────────────────────────────
    _add_bold_label_value(doc, "Background of the Speaker(s): ", report.speaker_background or "N/A")
    doc.add_paragraph()

    # ── REPORT ON THE SESSION ─────────────────────────────────────────────────
    p_rep = doc.add_paragraph()
    p_rep.add_run("Report on the Session:\n").bold = True
    p_rep.add_run("Session Summary:\n").bold = True
    p_rep.add_run(report.event_summary or "N/A")
    p_rep.add_run("\n\nDetailed Session Report:\n").bold = True
    p_rep.add_run(report.session_report or "N/A")

    if report.issues:
        p_rep.add_run("\n\nIssues Faced:\n").bold = True
        p_rep.add_run(report.issues)
    if report.feedback:
        p_rep.add_run("\n\nFeedback and Suggestions:\n").bold = True
        p_rep.add_run(report.feedback)
    doc.add_paragraph()

    # ── KEY OUTCOMES ──────────────────────────────────────────────────────────
    p_outcomes = doc.add_paragraph()
    p_outcomes.add_run("Key Outcomes:\n").bold = True
    key_outcomes = report.key_outcomes or []
    if key_outcomes:
        for outcome in key_outcomes:
            bp = doc.add_paragraph(style="List Bullet")
            bp.add_run(outcome)
    else:
        doc.add_paragraph("N/A")
    doc.add_paragraph()

    # ── CONCLUSION ────────────────────────────────────────────────────────────
    _add_bold_label_value(doc, "Conclusion: ", report.conclusion or "N/A")
    doc.add_paragraph()

    # ── ATTACHMENTS & GLIMPSES ────────────────────────────────────────────────
    p_glimpses = doc.add_paragraph()
    p_glimpses.add_run("Glimpses of the Event:\n").bold = True

    if report.attendance_doc_path:
        _add_bold_label_value(doc, "Attendance Document: ", os.path.basename(report.attendance_doc_path))

    # Flier (1 compulsory)
    if report.flier_path and os.path.exists(report.flier_path):
        _add_subsection_heading(doc, "Event Flier")
        try:
            if preserve_ratio:
                with PILImage.open(report.flier_path) as img:
                    w_px, h_px = img.size
                w_in = w_px / 96.0
                h_in = h_px / 96.0
                max_w = 4.0
                max_h = 3.0
                if w_in > max_w or h_in > max_h:
                    ratio = min(max_w / w_in, max_h / h_in)
                    w_in = w_in * ratio
                    h_in = h_in * ratio
                doc.add_picture(report.flier_path, width=Inches(w_in), height=Inches(h_in))
            else:
                doc.add_picture(report.flier_path, width=Inches(4.0), height=Inches(3.0))
            fp = doc.add_paragraph(os.path.basename(report.flier_path))
            fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        except Exception:
            doc.add_paragraph(f"Flier: {os.path.basename(report.flier_path)}")

    # Event photos (4–8, embed up to 8)
    photo_dir = os.path.join(settings.STORAGE_ROOT, "events", str(event.id), "rnd_report", "photos")
    if os.path.exists(photo_dir):
        photos = sorted([
            f for f in os.listdir(photo_dir)
            if f.lower().endswith((".jpg", ".jpeg", ".png", ".webp"))
        ])
        if photos:
            _add_subsection_heading(doc, "Event Photos")
            for photo in photos[:8]:
                photo_path = os.path.join(photo_dir, photo)
                try:
                    if preserve_ratio:
                        with PILImage.open(photo_path) as img:
                            w_px, h_px = img.size
                        w_in = w_px / 96.0
                        h_in = h_px / 96.0
                        max_w = 3.0
                        max_h = 2.0
                        if w_in > max_w or h_in > max_h:
                            ratio = min(max_w / w_in, max_h / h_in)
                            w_in = w_in * ratio
                            h_in = h_in * ratio
                        doc.add_picture(photo_path, width=Inches(w_in), height=Inches(h_in))
                    else:
                        doc.add_picture(photo_path, width=Inches(3.0), height=Inches(2.0))
                    pp = doc.add_paragraph(photo)
                    pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                except Exception:
                    doc.add_paragraph(f"Photo: {photo}")

    # ── FOOTER NOTE ────────────────────────────────────────────────────────────
    doc.add_paragraph()
    footer_note = doc.add_paragraph(
        f"RnD Report generated on {datetime.now().strftime('%d %B %Y at %I:%M %p')}"
    )
    footer_note.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    # ── SAVE ───────────────────────────────────────────────────────────────────
    report_dir = os.path.join(settings.STORAGE_ROOT, "events", str(event.id), "rnd_report")
    os.makedirs(report_dir, exist_ok=True)
    output_path = os.path.join(report_dir, "generated_rnd_report.docx")
    doc.save(output_path)

    return output_path
